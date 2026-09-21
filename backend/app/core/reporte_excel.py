from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_PRINCIPAL = "087F8C"
FORMATO_PESOS = '"$"#,##0'


def _formato_tabla(hoja, anchos):
    """Encabezado con color, columnas con ancho, fila fija y filtros activados."""
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_PRINCIPAL)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions


def generar_excel_reporte(reporte: dict) -> bytes:
    libro = Workbook()

    # ---------- Hoja 1: una fila por venta ----------
    ventas = libro.active
    ventas.title = "Ventas"
    ventas.append(["Fecha", "Hora", "N° venta", "Cliente", "Total", "Estado"])
    for v in reporte["ventas"]:
        ventas.append(
            [
                v["fecha_venta"].date(),
                v["fecha_venta"].time().replace(microsecond=0),
                v["numero_venta"],
                v["cliente"],
                float(v["total"]),
                v["estado"],
            ]
        )
    for fila in ventas.iter_rows(min_row=2):
        fila[0].number_format = "DD/MM/YYYY"
        fila[1].number_format = "HH:MM"
        fila[4].number_format = FORMATO_PESOS
    _formato_tabla(ventas, [13, 9, 14, 30, 16, 14])

    # ---------- Hoja 2: una fila por servicio vendido ----------
    detalle = libro.create_sheet("Detalle")
    detalle.append(
        ["Fecha", "N° venta", "Cliente", "Servicio", "Cantidad", "Precio unitario", "Subtotal", "Estado"]
    )
    for v in reporte["ventas"]:
        for item in v["items"]:
            detalle.append(
                [
                    v["fecha_venta"].date(),
                    v["numero_venta"],
                    v["cliente"],
                    item["servicio"],
                    item["cantidad"],
                    float(item["precio_unitario"]),
                    float(item["subtotal"]),
                    v["estado"],
                ]
            )
    for fila in detalle.iter_rows(min_row=2):
        fila[0].number_format = "DD/MM/YYYY"
        fila[5].number_format = FORMATO_PESOS
        fila[6].number_format = FORMATO_PESOS
    _formato_tabla(detalle, [13, 14, 30, 34, 10, 16, 16, 14])

    # ---------- Hoja 3: resumen del día ----------
    resumen = libro.create_sheet("Resumen")
    filas = [
        ("Fecha del reporte", reporte["fecha"]),
        ("Total vendido (ventas completadas)", float(reporte["total_vendido"])),
        ("Ventas del día", reporte["cantidad_ventas"]),
        ("Completadas", reporte["cantidad_completadas"]),
        ("Pendientes", reporte["cantidad_pendientes"]),
        ("Anuladas", reporte["cantidad_anuladas"]),
        ("Generado por", reporte["generado_por"]),
        ("Generado el", reporte["generado_en"]),
    ]
    for etiqueta, valor in filas:
        resumen.append([etiqueta, valor])
    for fila in resumen.iter_rows():
        fila[0].font = Font(bold=True)
        fila[1].alignment = Alignment(horizontal="left")
    resumen["B1"].number_format = "DD/MM/YYYY"
    resumen["B2"].number_format = FORMATO_PESOS
    resumen["B8"].number_format = "DD/MM/YYYY HH:MM"
    resumen.column_dimensions["A"].width = 38
    resumen.column_dimensions["B"].width = 24

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()